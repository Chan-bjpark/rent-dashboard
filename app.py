#!/usr/bin/env python3
"""월세 시세 분석 대시보드 — 국토부 실거래가 기반"""

import io
import os, math, json, time, re
import requests
import xml.etree.ElementTree as ET
from flask import Flask, render_template, jsonify, request as req, send_file
from datetime import datetime, timedelta
from collections import defaultdict
from dotenv import load_dotenv
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

load_dotenv()

app = Flask(__name__)

# ── Config ────────────────────────────────────────────────────────
DATA_KEY = os.environ.get("DATA_GO_KR_KEY", "")
KAKAO_REST = os.environ.get("KAKAO_REST_KEY", "")
KAKAO_JS = os.environ.get("KAKAO_JS_KEY", "")
PORT = int(os.environ.get("PORT", 8585))

# ── 국토부 API endpoints ─────────────────────────────────────────
MOLIT = {
    "officetel": {
        "url": "https://apis.data.go.kr/1613000/RTMSDataSvcOffiRent/getRTMSDataSvcOffiRent",
        "label": "오피스텔",
        "name_field": "offiNm",
    },
    "multi_family": {
        "url": "https://apis.data.go.kr/1613000/RTMSDataSvcRHRent/getRTMSDataSvcRHRent",
        "label": "연립다세대",
        "name_field": "mhouseNm",
    },
    "apartment": {
        "url": "https://apis.data.go.kr/1613000/RTMSDataSvcAptRent/getRTMSDataSvcAptRent",
        "label": "아파트",
        "name_field": "aptNm",
    },
}

# ── Cache ─────────────────────────────────────────────────────────
_cache = {}
CACHE_TTL = 3600


def cached(key, fn):
    now = time.time()
    if key in _cache and now - _cache[key]["t"] < CACHE_TTL:
        return _cache[key]["v"]
    val = fn()
    _cache[key] = {"v": val, "t": now}
    return val


# ── Utility ───────────────────────────────────────────────────────
def haversine(lat1, lon1, lat2, lon2):
    R = 6371
    dlat, dlon = math.radians(lat2 - lat1), math.radians(lon2 - lon1)
    a = math.sin(dlat / 2) ** 2 + math.cos(math.radians(lat1)) * math.cos(
        math.radians(lat2)
    ) * math.sin(dlon / 2) ** 2
    return R * 2 * math.asin(math.sqrt(a))


def clean_num(s):
    if not s:
        return 0
    return int(str(s).replace(",", "").strip() or 0)


# ── Kakao API ─────────────────────────────────────────────────────
def kakao_get(url, params):
    h = {"Authorization": f"KakaoAK {KAKAO_REST}"}
    r = requests.get(url, headers=h, params=params, timeout=5)
    r.raise_for_status()
    return r.json()


def kakao_search(query):
    """키워드/주소 검색 → 좌표"""
    try:
        data = kakao_get(
            "https://dapi.kakao.com/v2/local/search/keyword.json",
            {"query": query, "size": 5},
        )
        results = data.get("documents", [])
        if not results:
            data = kakao_get(
                "https://dapi.kakao.com/v2/local/search/address.json",
                {"query": query, "size": 5},
            )
            results = data.get("documents", [])
        return results
    except Exception as e:
        print(f"Kakao search error: {e}")
        return []


def kakao_coord2region(lng, lat):
    """좌표 → 법정동 정보"""
    try:
        data = kakao_get(
            "https://dapi.kakao.com/v2/local/geo/coord2regioncode.json",
            {"x": lng, "y": lat},
        )
        for doc in data.get("documents", []):
            if doc.get("region_type") == "B":
                return doc
        return data.get("documents", [{}])[0] if data.get("documents") else None
    except Exception as e:
        print(f"Kakao coord2region error: {e}")
        return None


def geocode_dong(sido, sigungu, dong):
    """법정동명 → 중심 좌표"""
    key = f"geo:{sido}:{sigungu}:{dong}"

    def _f():
        try:
            q = f"{sido} {sigungu} {dong}"
            data = kakao_get(
                "https://dapi.kakao.com/v2/local/search/address.json",
                {"query": q, "size": 1},
            )
            docs = data.get("documents", [])
            if docs:
                return {"lat": float(docs[0]["y"]), "lng": float(docs[0]["x"])}
            # fallback keyword
            data = kakao_get(
                "https://dapi.kakao.com/v2/local/search/keyword.json",
                {"query": q, "size": 1},
            )
            docs = data.get("documents", [])
            if docs:
                return {"lat": float(docs[0]["y"]), "lng": float(docs[0]["x"])}
        except Exception:
            pass
        return None

    return cached(key, _f)


# ── 국토부 API ────────────────────────────────────────────────────
def fetch_molit(ptype, region_code, deal_ym):
    api = MOLIT.get(ptype)
    if not api:
        return []
    ck = f"molit:{ptype}:{region_code}:{deal_ym}"

    def _f():
        try:
            r = requests.get(
                api["url"],
                params={
                    "serviceKey": DATA_KEY,
                    "LAWD_CD": region_code,
                    "DEAL_YMD": deal_ym,
                    "numOfRows": 9999,
                    "pageNo": 1,
                },
                headers={"User-Agent": "Mozilla/5.0"},
                timeout=15,
            )
            r.raise_for_status()
            return parse_xml(r.text, api)
        except Exception as e:
            print(f"MOLIT API error [{ptype} {region_code} {deal_ym}]: {e}")
            return []

    return cached(ck, _f)


def parse_xml(xml_text, api_info):
    items = []
    try:
        root = ET.fromstring(xml_text)
        # 에러 체크
        code_el = root.find(".//resultCode")
        if code_el is not None and code_el.text not in ("00", "000"):
            msg = root.findtext(".//resultMsg", "")
            print(f"API returned error: {code_el.text} {msg}")
            return []

        for item in root.iter("item"):
            g = lambda tag: (item.findtext(tag) or "").strip()
            deposit = clean_num(g("deposit") or g("보증금액"))
            monthly = clean_num(g("monthlyRent") or g("월세금액"))
            area_str = g("excluUseAr") or g("전용면적") or "0"
            floor_str = g("floor") or g("층") or "0"
            built_str = g("buildYear") or g("건축년도") or "0"

            items.append(
                {
                    "type": api_info["label"],
                    "building": g(api_info["name_field"]),
                    "dong": g("umdNm") or g("법정동"),
                    "jibun": g("jibun") or g("지번"),
                    "area": float(area_str) if area_str else 0,
                    "floor": int(floor_str) if floor_str.lstrip("-").isdigit() else 0,
                    "built_year": int(built_str) if built_str.isdigit() else 0,
                    "deposit": deposit,
                    "monthly_rent": monthly,
                    "year": g("dealYear") or g("년"),
                    "month": g("dealMonth") or g("월"),
                    "day": g("dealDay") or g("일"),
                    "contract_type": "월세" if monthly > 0 else "전세",
                }
            )
    except ET.ParseError as e:
        print(f"XML parse error: {e}")
    return items


# ── Statistics ────────────────────────────────────────────────────
def calc_stats(data):
    if not data:
        return {"count": 0}

    rents = sorted([d["monthly_rent"] for d in data if d["monthly_rent"] > 0])
    deps = sorted([d["deposit"] for d in data])
    areas = [d["area"] for d in data if d["area"] > 0]

    def quartiles(arr):
        if not arr:
            return {}
        n = len(arr)
        return {
            "mean": round(sum(arr) / n),
            "median": arr[n // 2],
            "min": arr[0],
            "max": arr[-1],
            "q1": arr[n // 4],
            "q3": arr[3 * n // 4] if n > 1 else arr[0],
            "count": n,
        }

    stats = {
        "count": len(data),
        "rent": quartiles(rents),
        "deposit": quartiles(deps),
    }

    # 평당 월세
    per_pyeong = []
    for d in data:
        if d["monthly_rent"] > 0 and d["area"] > 0:
            per_pyeong.append(d["monthly_rent"] / (d["area"] / 3.3058))
    if per_pyeong:
        per_pyeong.sort()
        n = len(per_pyeong)
        stats["rent_per_pyeong"] = {
            "mean": round(sum(per_pyeong) / n, 1),
            "median": round(per_pyeong[n // 2], 1),
        }

    # 건물유형별
    by_type = defaultdict(list)
    for d in data:
        if d["monthly_rent"] > 0:
            by_type[d["type"]].append(d["monthly_rent"])
    stats["by_type"] = {}
    for t, rs in by_type.items():
        rs.sort()
        n = len(rs)
        stats["by_type"][t] = {
            "count": n,
            "mean": round(sum(rs) / n),
            "median": rs[n // 2],
        }

    # 면적구간별
    ranges = [(0, 30, "30㎡ 미만"), (30, 999, "30㎡ 이상")]
    stats["by_area"] = {}
    for lo, hi, label in ranges:
        rs = sorted([d["monthly_rent"] for d in data if d["monthly_rent"] > 0 and lo <= d["area"] < hi])
        if rs:
            n = len(rs)
            stats["by_area"][label] = {"count": n, "mean": round(sum(rs) / n), "median": rs[n // 2]}

    # 층수별
    stats["by_floor"] = {}
    for lo, hi, label in [(1, 5, "저층(1~4)"), (5, 11, "중층(5~10)"), (11, 21, "고층(11~20)"), (21, 999, "초고층(21+)")]:
        rs = sorted([d["monthly_rent"] for d in data if d["monthly_rent"] > 0 and lo <= d["floor"] < hi])
        if rs:
            n = len(rs)
            stats["by_floor"][label] = {"count": n, "mean": round(sum(rs) / n), "median": rs[n // 2]}

    return stats


# ── Routes ────────────────────────────────────────────────────────
@app.route("/")
def index():
    return render_template("index.html", kakao_js_key=KAKAO_JS)


@app.route("/api/search")
def api_search():
    q = req.args.get("q", "")
    if not q:
        return jsonify({"results": []})
    docs = kakao_search(q)
    results = []
    for d in docs:
        results.append(
            {
                "name": d.get("place_name") or d.get("address_name", ""),
                "address": d.get("address_name") or d.get("road_address_name", ""),
                "lat": float(d.get("y", 0)),
                "lng": float(d.get("x", 0)),
            }
        )
    return jsonify({"results": results})


@app.route("/api/rent")
def api_rent():
    lat = float(req.args.get("lat", 0))
    lng = float(req.args.get("lng", 0))
    radius = float(req.args.get("radius", 1.0))
    ptypes = req.args.get("types", "officetel,multi_family").split(",")
    months = int(req.args.get("months", 6))

    if not lat or not lng:
        return jsonify({"error": "좌표가 필요합니다"}), 400

    # 1) 좌표 → 시군구 코드
    region = kakao_coord2region(lng, lat)
    if not region:
        return jsonify({"error": "지역 정보를 찾을 수 없습니다"}), 404

    code = region.get("code", "")[:5]
    sido = region.get("region_1depth_name", "")
    sigungu = region.get("region_2depth_name", "")

    # 인접 시군구 체크 (반경 경계)
    codes = {code}
    for bearing in [0, 90, 180, 270]:
        dlat = radius / 111.0 * math.cos(math.radians(bearing))
        dlng = radius / (111.0 * math.cos(math.radians(lat))) * math.sin(math.radians(bearing))
        br = kakao_coord2region(lng + dlng, lat + dlat)
        if br:
            codes.add(br.get("code", "")[:5])
    codes.discard("")

    # 2) 최근 N개월
    yms = set()
    now = datetime.now()
    for i in range(months):
        dt = now - timedelta(days=30 * i)
        yms.add(dt.strftime("%Y%m"))

    # 3) API 호출
    all_data = []
    for c in codes:
        for pt in ptypes:
            pt = pt.strip()
            if pt not in MOLIT:
                continue
            for ym in sorted(yms, reverse=True):
                all_data.extend(fetch_molit(pt, c, ym))

    # 4) 법정동별 좌표 → 거리 필터
    dong_coords = {}
    unique_dongs = {d["dong"] for d in all_data if d["dong"]}
    for dong in unique_dongs:
        coords = geocode_dong(sido, sigungu, dong)
        if coords:
            dong_coords[dong] = coords

    filtered = []
    for rec in all_data:
        d = rec["dong"]
        if d in dong_coords:
            c = dong_coords[d]
            dist = haversine(lat, lng, c["lat"], c["lng"])
            rec["distance"] = round(dist, 2)
            rec["dong_lat"] = c["lat"]
            rec["dong_lng"] = c["lng"]
            if dist <= radius:
                filtered.append(rec)
        else:
            rec["distance"] = None
            rec["dong_lat"] = None
            rec["dong_lng"] = None
            filtered.append(rec)

    stats = calc_stats(filtered)

    return jsonify(
        {
            "data": filtered,
            "stats": stats,
            "region": {"code": code, "sido": sido, "sigungu": sigungu},
            "dong_coords": dong_coords,
            "total": len(filtered),
        }
    )


# ── 시군구 코드 직접 조회 (Kakao 없을 때 fallback) ────────────────
REGION_CODES = {
    "서울 종로구": "11110", "서울 중구": "11140", "서울 용산구": "11170",
    "서울 성동구": "11200", "서울 광진구": "11215", "서울 동대문구": "11230",
    "서울 중랑구": "11260", "서울 성북구": "11290", "서울 강북구": "11305",
    "서울 도봉구": "11320", "서울 노원구": "11350", "서울 은평구": "11380",
    "서울 서대문구": "11410", "서울 마포구": "11440", "서울 양천구": "11470",
    "서울 강서구": "11500", "서울 구로구": "11530", "서울 금천구": "11545",
    "서울 영등포구": "11560", "서울 동작구": "11590", "서울 관악구": "11620",
    "서울 서초구": "11650", "서울 강남구": "11680", "서울 송파구": "11710",
    "서울 강동구": "11740",
    "경기 수원 장안구": "41111", "경기 수원 권선구": "41113",
    "경기 수원 팔달구": "41115", "경기 수원 영통구": "41117",
    "경기 성남 수정구": "41131", "경기 성남 중원구": "41133",
    "경기 성남 분당구": "41135", "경기 의정부시": "41150",
    "경기 안양 만안구": "41171", "경기 안양 동안구": "41173",
    "경기 부천시": "41190", "경기 광명시": "41210",
    "경기 평택시": "41220", "경기 동두천시": "41250",
    "경기 안산 상록구": "41271", "경기 안산 단원구": "41273",
    "경기 고양 덕양구": "41281", "경기 고양 일산동구": "41285",
    "경기 고양 일산서구": "41287", "경기 과천시": "41290",
    "경기 구리시": "41310", "경기 남양주시": "41360",
    "경기 오산시": "41370", "경기 시흥시": "41390",
    "경기 군포시": "41410", "경기 의왕시": "41430",
    "경기 하남시": "41450", "경기 용인 처인구": "41461",
    "경기 용인 기흥구": "41463", "경기 용인 수지구": "41465",
    "경기 파주시": "41480", "경기 이천시": "41500",
    "경기 안성시": "41550", "경기 김포시": "41570",
    "경기 화성시": "41590", "경기 광주시": "41610",
    "경기 양주시": "41630", "경기 포천시": "41650",
    "부산 중구": "26110", "부산 서구": "26140", "부산 동구": "26170",
    "부산 영도구": "26200", "부산 부산진구": "26230", "부산 동래구": "26260",
    "부산 남구": "26290", "부산 북구": "26320", "부산 해운대구": "26350",
    "부산 사하구": "26380", "부산 금정구": "26410", "부산 강서구": "26440",
    "부산 연제구": "26470", "부산 수영구": "26500", "부산 사상구": "26530",
    "부산 기장군": "26710",
    "인천 중구": "28110", "인천 동구": "28140", "인천 미추홀구": "28177",
    "인천 연수구": "28185", "인천 남동구": "28200", "인천 부평구": "28237",
    "인천 계양구": "28245", "인천 서구": "28260",
    "대구 중구": "27110", "대구 동구": "27140", "대구 서구": "27170",
    "대구 남구": "27200", "대구 북구": "27230", "대구 수성구": "27260",
    "대구 달서구": "27290",
    "대전 동구": "30110", "대전 중구": "30140", "대전 서구": "30170",
    "대전 유성구": "30200", "대전 대덕구": "30230",
    "광주 동구": "29110", "광주 서구": "29140", "광주 남구": "29155",
    "광주 북구": "29170", "광주 광산구": "29200",
}


@app.route("/api/rent_direct")
def api_rent_direct():
    """Kakao API 없이 시군구 코드 직접 조회"""
    code = req.args.get("code", "")
    ptypes = req.args.get("types", "officetel,multi_family").split(",")
    months = int(req.args.get("months", 6))

    if not code:
        return jsonify({"error": "시군구 코드 필요"}), 400

    yms = set()
    now = datetime.now()
    for i in range(months):
        dt = now - timedelta(days=30 * i)
        yms.add(dt.strftime("%Y%m"))

    all_data = []
    for pt in ptypes:
        pt = pt.strip()
        if pt not in MOLIT:
            continue
        for ym in sorted(yms, reverse=True):
            all_data.extend(fetch_molit(pt, code, ym))

    # 월세만 필터
    monthly_only = [d for d in all_data if d["monthly_rent"] > 0]

    stats = calc_stats(monthly_only)
    return jsonify({"data": all_data, "stats": stats, "total": len(all_data)})


@app.route("/api/regions")
def api_regions():
    return jsonify(REGION_CODES)


# ── Export: Google Slides (PPTX) ─────────────────────────────────
NAVY = RGBColor(0x1B, 0x20, 0x38)
BLUE = RGBColor(0x3B, 0x5F, 0xCC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x7F, 0x8C, 0x8D)
GREEN = RGBColor(0x2E, 0xA0, 0x43)
RED = RGBColor(0xE7, 0x4C, 0x3C)
LIGHT_BLUE_BG = RGBColor(0xE8, 0xF0, 0xFE)
BEIGE = RGBColor(0xE8, 0xE0, 0xD4)


def _add_header_bar(slide, title_text):
    """네이비 헤더 바 + 파란 액센트 + 제목"""
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Inches(0.7))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    accent = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.06), Inches(0.7))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(10), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE


def _add_kpi_card(slide, left, top, width, height, label, value, sub="", color=BLUE):
    """KPI 카드 (라벨 + 큰 숫자 + 서브텍스트)"""
    box = slide.shapes.add_shape(1, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    box.line.width = Pt(0.5)
    txBox = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), width - Inches(0.3), height - Inches(0.2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(9)
    p.font.color.rgb = GRAY
    p2 = tf.add_paragraph()
    p2.text = str(value)
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = color
    p2.space_before = Pt(4)
    if sub:
        p3 = tf.add_paragraph()
        p3.text = sub
        p3.font.size = Pt(8)
        p3.font.color.rgb = GRAY


def _add_table(slide, left, top, width, headers, rows, title=""):
    """테이블 추가"""
    if title:
        txBox = slide.shapes.add_textbox(left, top - Inches(0.3), width, Inches(0.3))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = DARK

    n_rows = min(len(rows) + 1, 30)
    n_cols = len(headers)
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, Inches(0.25 * n_rows))
    tbl = tbl_shape.table

    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for para in cell.text_frame.paragraphs:
            para.font.size = Pt(8)
            para.font.bold = True
            para.font.color.rgb = WHITE
            para.alignment = PP_ALIGN.CENTER

    for r_idx, row in enumerate(rows[:n_rows - 1]):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.text = str(val)
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(8)
                para.font.color.rgb = DARK
                para.alignment = PP_ALIGN.CENTER


@app.route("/api/export/slides", methods=["POST"])
def export_slides():
    body = req.json or {}
    stats = body.get("stats", {})
    breakdowns = body.get("breakdowns", {})
    data = body.get("data", [])
    location = body.get("location", "")
    filters = body.get("filters", "")

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ── Slide 1: 타이틀 ──
    sl = prs.slides.add_slide(blank)
    bg = sl.background.fill
    bg.solid()
    bg.fore_color.rgb = BEIGE
    txBox = sl.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "월세 시세 분석 보고서"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    p2.text = f"국토교통부 실거래가 기반"
    p2.font.size = Pt(16)
    p2.font.color.rgb = GRAY
    p2.space_before = Pt(12)
    if location:
        p3 = tf.add_paragraph()
        p3.text = location
        p3.font.size = Pt(14)
        p3.font.color.rgb = BLUE
        p3.space_before = Pt(8)
    txDate = sl.shapes.add_textbox(Inches(9.5), Inches(6), Inches(3), Inches(0.5))
    tf2 = txDate.text_frame
    p = tf2.paragraphs[0]
    p.text = datetime.now().strftime("%Y.%m.%d")
    p.font.size = Pt(12)
    p.font.color.rgb = GRAY
    p.alignment = PP_ALIGN.RIGHT

    # ── Slide 2: 시세 요약 ──
    sl = prs.slides.add_slide(blank)
    _add_header_bar(sl, "시세 요약")
    rent = stats.get("rent", {})
    dep = stats.get("deposit", {})
    pp = stats.get("rent_per_pyeong", {})

    cards = [
        ("거래 건수", f"{stats.get('count', 0)}건", "", BLUE),
        ("평균 월세", f"{rent.get('mean', '-')}만원", f"중위 {rent.get('median', '-')}만원", RGBColor(0xD2, 0x99, 0x22)),
        ("월세 범위", f"{rent.get('min', 0)} ~ {rent.get('max', 0)}만원", f"Q1:{rent.get('q1', '-')} Q3:{rent.get('q3', '-')}", DARK),
        ("평균 보증금", f"{dep.get('mean', '-')}만원", f"중위 {dep.get('median', '-')}만원", GREEN),
        ("보증금 범위", f"{dep.get('min', 0)} ~ {dep.get('max', 0)}만원", "", DARK),
        ("평당 월세", f"{pp.get('median', '-')}만원", f"평균 {pp.get('mean', '-')}만/평", RGBColor(0xBC, 0x8C, 0xFF)),
    ]
    for i, (label, value, sub, color) in enumerate(cards):
        col = i % 3
        row = i // 3
        _add_kpi_card(sl, Inches(0.5 + col * 4.2), Inches(1.2 + row * 2.8), Inches(3.8), Inches(2.2), label, value, sub, color)

    if filters:
        txBox = sl.shapes.add_textbox(Inches(0.5), Inches(6.8), Inches(12), Inches(0.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = f"조건: {filters}"
        p.font.size = Pt(8)
        p.font.color.rgb = GRAY

    # ── Slide 3: 월세 분포 차트 ──
    sl = prs.slides.add_slide(blank)
    _add_header_bar(sl, "월세 분포 차트")

    for ci, (chart_label, chart_data) in enumerate([
        ("30㎡ 미만 월세 분포", breakdowns.get("under30_hist", {})),
        ("30㎡ 이상 월세 분포", breakdowns.get("over30_hist", {})),
    ]):
        labels = chart_data.get("labels", [])
        values = chart_data.get("values", [])
        if not labels:
            continue
        chart_d = slide_chart_data(labels, values)
        left = Inches(0.5 + ci * 6.3)
        x, y, cx, cy = left, Inches(1.2), Inches(5.8), Inches(5.5)
        chart_frame = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_d)
        chart = chart_frame.chart
        chart.has_legend = False
        plot = chart.plots[0]
        series = plot.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = BLUE if ci == 0 else RGBColor(0xBC, 0x8C, 0xFF)
        chart.chart_title.has_text_frame = True
        chart.chart_title.text_frame.paragraphs[0].text = f"{chart_label} (만원)"
        chart.chart_title.text_frame.paragraphs[0].font.size = Pt(11)
        chart.chart_title.text_frame.paragraphs[0].font.color.rgb = DARK

    # ── Slide 4: 조건별 시세 비교 ──
    sl = prs.slides.add_slide(blank)
    _add_header_bar(sl, "조건별 시세 비교")

    tbl_headers = ["구분", "건수", "평균(만원)", "중위(만원)"]
    tables = [
        ("건물유형별", breakdowns.get("by_type", [])),
        ("면적구간별", breakdowns.get("by_area", [])),
        ("층수별", breakdowns.get("by_floor", [])),
        ("법정동별", breakdowns.get("by_dong", [])),
    ]
    for i, (title, rows) in enumerate(tables):
        col = i % 2
        row = i // 2
        left = Inches(0.5 + col * 6.3)
        top = Inches(1.4 + row * 3.0)
        _add_table(sl, left, top, Inches(5.8), tbl_headers, rows, title)

    # ── Slide 5: 거래 내역 (상위 25건) ──
    sl = prs.slides.add_slide(blank)
    _add_header_bar(sl, f"거래 내역 (총 {len(data)}건 중 상위 25건)")

    tx_headers = ["유형", "건물명", "법정동", "면적(㎡)", "층", "보증금(만)", "월세(만)", "건축", "계약일", "구분"]
    tx_rows = []
    for d in data[:25]:
        tx_rows.append([
            d.get("type", ""), d.get("building", "-"), d.get("dong", ""),
            str(d.get("area", "")), str(d.get("floor", "")),
            f"{d.get('deposit', 0):,}", f"{d.get('monthly_rent', 0):,}",
            str(d.get("built_year", "-")),
            f"{d.get('year', '')}.{str(d.get('month', '')).zfill(2)}.{str(d.get('day', '')).zfill(2)}",
            d.get("contract_type", ""),
        ])
    _add_table(sl, Inches(0.3), Inches(1.0), Inches(12.7), tx_headers, tx_rows)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    fname = f"월세시세보고서_{datetime.now().strftime('%Y%m%d')}.pptx"
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")


def slide_chart_data(labels, values):
    from pptx.chart.data import CategoryChartData
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series("건수", values)
    return cd


# ── Export: Excel ────────────────────────────────────────────────
@app.route("/api/export/excel", methods=["POST"])
def export_excel():
    body = req.json or {}
    data = body.get("data", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "거래내역"

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="1B2038", end_color="1B2038", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    headers = ["유형", "건물명", "법정동", "지번", "전용면적(㎡)", "층",
               "보증금(만원)", "월세(만원)", "건축년도", "계약일", "구분"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for r, d in enumerate(data, 2):
        vals = [
            d.get("type", ""), d.get("building", ""), d.get("dong", ""),
            d.get("jibun", ""), d.get("area", 0), d.get("floor", 0),
            d.get("deposit", 0), d.get("monthly_rent", 0),
            d.get("built_year", ""),
            f"{d.get('year', '')}.{str(d.get('month', '')).zfill(2)}.{str(d.get('day', '')).zfill(2)}",
            d.get("contract_type", ""),
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = thin_border
            if isinstance(v, (int, float)) and c in (5, 6, 7, 8):
                cell.number_format = '#,##0'

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + c) if c <= 26 else 'A'].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"거래내역_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    print(f"🏠 월세 시세 대시보드: http://localhost:{PORT}")
    print(f"   DATA_GO_KR_KEY: {'✓' if DATA_KEY else '✗'}")
    print(f"   KAKAO_REST_KEY: {'✓' if KAKAO_REST else '✗ (지도/검색 제한)'}")
    print(f"   KAKAO_JS_KEY:   {'✓' if KAKAO_JS else '✗ (지도 표시 제한)'}")
    app.run(host="0.0.0.0", port=PORT, debug=True)
